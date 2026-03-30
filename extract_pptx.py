from pptx import Presentation

# Open the PowerPoint file
prs = Presentation('BTP Presentation (4).pptx')

print("=" * 80)
print("POWERPOINT CONTENT EXTRACTION")
print("=" * 80)

for slide_num, slide in enumerate(prs.slides, 1):
    print(f"\n{'='*80}")
    print(f"SLIDE {slide_num}")
    print(f"{'='*80}")
    
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text.strip():
            print(shape.text)
            print()

print("\n" + "=" * 80)
print("END OF PRESENTATION")
print("=" * 80)
